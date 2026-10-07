"""Modul loader dan normalisasi dokumen lokal.

Mendukung format PDF, DOCX, XLSX, MD, dan TXT menjadi representasi
blok terstruktur (Block) dengan metadata lokasi presisi.
"""

import argparse
from dataclasses import asdict, dataclass
from datetime import datetime
import hashlib
import logging
from pathlib import Path
import re
from typing import Any, Optional

import docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
import openpyxl
import pymupdf

from app.config import Settings, get_settings
from app.logging_setup import setup_logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".md", ".txt"}


@dataclass
class Block:
    """Representasi satu unit teks atau tabel dari dokumen."""

    rel_path: str
    file_hash: str
    filename: str
    file_type: str
    folder: str
    modified_at: str
    page: Optional[int] = None
    sheet: Optional[str] = None
    section: Optional[str] = None
    kind: str = "paragraph"  # "heading", "paragraph", "table", "list", "code"
    text: str = ""
    needs_ocr: bool = False

    def to_dict(self) -> dict[str, Any]:
        """Konversi block ke dictionary."""
        return asdict(self)


def format_location(block: Block) -> str:
    """Menghasilkan representasi lokasi yang presisi untuk sitasi."""
    parts = [block.rel_path]
    if block.sheet:
        parts.append(f"sheet: {block.sheet}")
    if block.page is not None:
        parts.append(f"hal. {block.page}")
    if block.section:
        parts.append(block.section)
    return f"[{' | '.join(parts)}]"


def compute_file_hash(path: Path) -> str:
    """Menghitung SHA-256 konten file dengan streaming 64KB."""
    sha256 = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            sha256.update(chunk)
    return sha256.hexdigest()


def get_file_metadata(path: Path, root_dir: Optional[Path] = None) -> dict[str, Any]:
    """Mengambil metadata standar file dokumen."""
    abs_path = path.resolve()
    if root_dir is not None:
        try:
            rel_path = str(abs_path.relative_to(root_dir.resolve()).as_posix())
        except ValueError:
            rel_path = str(path.as_posix())
    else:
        rel_path = str(path.as_posix())

    stat = path.stat()
    folder = ""
    if "/" in rel_path:
        folder = rel_path.rsplit("/", 1)[0]
    elif "\\" in rel_path:
        folder = rel_path.rsplit("\\", 1)[0]

    return {
        "rel_path": rel_path,
        "file_hash": compute_file_hash(path),
        "filename": path.name,
        "file_type": path.suffix.lstrip(".").lower(),
        "folder": folder,
        "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }


def scan_directory(root_dir: Path, max_file_mb: int = 50) -> list[Path]:
    """Memindai direktori dokumen dan menyaring file yang didukung."""
    if not root_dir.exists() or not root_dir.is_dir():
        logger.warning(f"Direktori tidak ditemukan: {root_dir}")
        return []

    found_files: list[Path] = []
    max_bytes = max_file_mb * 1024 * 1024

    for p in sorted(root_dir.rglob("*")):
        if not p.is_file():
            continue
        # Abaikan file/folder tersembunyi
        try:
            rel_parts = p.relative_to(root_dir).parts
            if any(part.startswith(".") for part in rel_parts):
                continue
        except ValueError:
            if p.name.startswith("."):
                continue

        # Abaikan file temporer Office
        if p.name.startswith("~$"):
            continue

        # Cek ekstensi yang didukung
        if p.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        # Cek batas ukuran file
        try:
            size = p.stat().st_size
            if size > max_bytes:
                logger.warning(
                    f"File dilewati karena melebihi batas {max_file_mb}MB: {p.name} ({size / (1024*1024):.1f}MB)"
                )
                continue
        except OSError as exc:
            logger.warning(f"Gagal memeriksa ukuran file {p.name}: {exc}")
            continue

        found_files.append(p)

    return found_files


def format_table_as_markdown(headers: list[str], rows: list[list[str]]) -> str:
    """Mengubah data matriks tabel menjadi representasi Markdown table."""
    if not headers and not rows:
        return ""
    if not headers and rows:
        headers = [f"Kolom {i + 1}" for i in range(len(rows[0]))]

    num_cols = max(len(headers), max((len(r) for r in rows), default=0))
    pad_headers = headers + [""] * (num_cols - len(headers))
    clean_headers = [str(h or "").strip().replace("\n", " ").replace("|", "/") for h in pad_headers]

    lines = [
        "| " + " | ".join(clean_headers) + " |",
        "| " + " | ".join(["---"] * num_cols) + " |",
    ]
    for r in rows:
        pad_row = r + [""] * (num_cols - len(r))
        clean_row = [str(c or "").strip().replace("\n", " ").replace("|", "/") for c in pad_row]
        lines.append("| " + " | ".join(clean_row) + " |")

    return "\n".join(lines)


# ==============================================================================
# Format Extractors
# ==============================================================================


def extract_pdf(path: Path, meta: dict[str, Any], settings: Settings) -> list[Block]:
    """Ekstraksi teks dan tabel dari dokumen PDF menggunakan PyMuPDF."""
    blocks: list[Block] = []
    doc = pymupdf.open(path)
    current_section: Optional[str] = None

    try:
        for page_index, page in enumerate(doc):
            page_num = page_index + 1
            tables_finder = page.find_tables()

            # Filter tabel bersarang (nested cell boxes)
            valid_tables = []
            for t in tables_finder.tables:
                t_rect = pymupdf.Rect(t.bbox)
                is_nested = any(
                    pymupdf.Rect(other.bbox).contains(t_rect)
                    and pymupdf.Rect(other.bbox).get_area() > t_rect.get_area()
                    for other in tables_finder.tables
                    if other is not t
                )
                if not is_nested:
                    valid_tables.append(t)

            table_rects = [pymupdf.Rect(t.bbox) for t in valid_tables]

            page_items: list[dict[str, Any]] = []

            # 1. Tambahkan tabel valid
            for t in valid_tables:
                extracted_rows = t.extract()
                if not extracted_rows:
                    continue
                headers = [str(c or "").strip() for c in extracted_rows[0]]
                data_rows = extracted_rows[1:]
                table_md = format_table_as_markdown(headers, data_rows)
                if table_md.strip():
                    page_items.append(
                        {
                            "y0": t.bbox[1],
                            "kind": "table",
                            "text": table_md,
                            "section": current_section,
                        }
                    )

            # 2. Tambahkan blok teks non-tabel
            text_blocks = page.get_text("blocks")
            for b in text_blocks:
                # b: (x0, y0, x1, y2, text, block_no, block_type)
                if b[6] != 0:
                    continue
                b_rect = pymupdf.Rect(b[:4])
                # Abaikan jika overlap signifikan (>50%) dengan tabel yang diekstrak
                is_in_table = any(
                    t_rect.intersects(b_rect)
                    and (t_rect & b_rect).get_area() / b_rect.get_area() > 0.5
                    for t_rect in table_rects
                )
                if is_in_table:
                    continue

                clean_text = b[4].strip()
                if not clean_text:
                    continue

                # Deteksi heading sederhana
                lines = [ln.strip() for ln in clean_text.splitlines() if ln.strip()]
                is_heading = False
                if len(lines) == 1 and len(clean_text) < 90:
                    first_line = lines[0]
                    if not first_line.endswith((".", ";", ":")) and (
                        first_line.isupper()
                        or first_line.startswith(("Bab ", "BAB ", "Pasal ", "PASAL ", "Tabel ", "TABEL "))
                        or first_line.istitle()
                    ):
                        is_heading = True
                        current_section = first_line

                page_items.append(
                    {
                        "y0": b[1],
                        "kind": "heading" if is_heading else "paragraph",
                        "text": clean_text,
                        "section": current_section,
                    }
                )

            # Urutkan berdasarkan koordinat vertikal (y0) agar urutan baca natural
            page_items.sort(key=lambda item: item["y0"])

            for item in page_items:
                blocks.append(
                    Block(
                        rel_path=meta["rel_path"],
                        file_hash=meta["file_hash"],
                        filename=meta["filename"],
                        file_type=meta["file_type"],
                        folder=meta["folder"],
                        modified_at=meta["modified_at"],
                        page=page_num,
                        sheet=None,
                        section=item["section"],
                        kind=item["kind"],
                        text=item["text"],
                        needs_ocr=False,
                    )
                )

            # 3. Cek halaman pindaian (scanned image) yang membutuhkan OCR
            full_page_text = page.get_text("text").strip()
            if len(full_page_text) < settings.OCR_MIN_CHARS:
                images = page.get_images()
                if len(images) > 0:
                    blocks.append(
                        Block(
                            rel_path=meta["rel_path"],
                            file_hash=meta["file_hash"],
                            filename=meta["filename"],
                            file_type=meta["file_type"],
                            folder=meta["folder"],
                            modified_at=meta["modified_at"],
                            page=page_num,
                            sheet=None,
                            section=current_section,
                            kind="paragraph",
                            text=f"[Halaman {page_num} berupa pindaian/gambar dan memerlukan OCR untuk diekstrak]",
                            needs_ocr=True,
                        )
                    )

    finally:
        doc.close()

    return blocks


def extract_docx(path: Path, meta: dict[str, Any], settings: Settings) -> list[Block]:
    """Ekstraksi teks sekuensial dan tabel dari dokumen DOCX."""
    blocks: list[Block] = []
    doc = docx.Document(path)
    current_section: Optional[str] = None

    for child in doc.element.body:
        # Elemen paragraf
        if child.tag.endswith("p"):
            p = DocxParagraph(child, doc)
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            is_heading = (
                style_name.startswith(("Heading", "heading", "Judul", "Title", "Subheading"))
                or (len(text) < 90 and text.startswith(("Pasal ", "PASAL ", "BAB ", "Bab ")))
            )

            if is_heading:
                current_section = text
                blocks.append(
                    Block(
                        rel_path=meta["rel_path"],
                        file_hash=meta["file_hash"],
                        filename=meta["filename"],
                        file_type=meta["file_type"],
                        folder=meta["folder"],
                        modified_at=meta["modified_at"],
                        page=None,
                        sheet=None,
                        section=current_section,
                        kind="heading",
                        text=text,
                    )
                )
            else:
                blocks.append(
                    Block(
                        rel_path=meta["rel_path"],
                        file_hash=meta["file_hash"],
                        filename=meta["filename"],
                        file_type=meta["file_type"],
                        folder=meta["folder"],
                        modified_at=meta["modified_at"],
                        page=None,
                        sheet=None,
                        section=current_section,
                        kind="paragraph",
                        text=text,
                    )
                )

        # Elemen tabel
        elif child.tag.endswith("tbl"):
            tbl = DocxTable(child, doc)
            if not tbl.rows:
                continue

            extracted_rows: list[list[str]] = []
            for row in tbl.rows:
                # Hilangkan sel duplikat dari merged cells berdampingan
                cell_texts: list[str] = []
                for cell in row.cells:
                    ct = cell.text.strip()
                    # python-docx mengembalikan sel yang sama beberapa kali untuk merged cells
                    if not cell_texts or ct != cell_texts[-1]:
                        cell_texts.append(ct)
                extracted_rows.append(cell_texts)

            if not extracted_rows:
                continue

            headers = extracted_rows[0]
            data_rows = extracted_rows[1:]
            table_md = format_table_as_markdown(headers, data_rows)
            if table_md.strip():
                blocks.append(
                    Block(
                        rel_path=meta["rel_path"],
                        file_hash=meta["file_hash"],
                        filename=meta["filename"],
                        file_type=meta["file_type"],
                        folder=meta["folder"],
                        modified_at=meta["modified_at"],
                        page=None,
                        sheet=None,
                        section=current_section,
                        kind="table",
                        text=table_md,
                    )
                )

    return blocks


def _format_cell_value(val: Any) -> str:
    """Format nilai sel Excel ke string representatif."""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S" if val.hour or val.minute else "%Y-%m-%d")
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.4f}".rstrip("0").rstrip(".")
    return str(val).strip()


def extract_xlsx(path: Path, meta: dict[str, Any], settings: Settings) -> list[Block]:
    """Ekstraksi data spreadsheet XLSX dengan chunking baris dan header berulang."""
    blocks: list[Block] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chunk_size = max(5, settings.XLSX_ROW_CHUNK)

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.sheet_state != "visible":
                continue

            rows_data: list[list[str]] = []
            for raw_row in ws.iter_rows(values_only=True):
                # Filter baris yang seluruh isinya None atau string kosong
                if any(c is not None and str(c).strip() != "" for c in raw_row):
                    formatted_row = [_format_cell_value(c) for c in raw_row]
                    rows_data.append(formatted_row)

            if not rows_data:
                continue

            headers = rows_data[0]
            body_rows = rows_data[1:]

            if not body_rows:
                # Hanya header tanpa data
                table_md = format_table_as_markdown(headers, [])
                blocks.append(
                    Block(
                        rel_path=meta["rel_path"],
                        file_hash=meta["file_hash"],
                        filename=meta["filename"],
                        file_type=meta["file_type"],
                        folder=meta["folder"],
                        modified_at=meta["modified_at"],
                        page=None,
                        sheet=sheet_name,
                        section=f"Sheet: {sheet_name}",
                        kind="table",
                        text=table_md,
                    )
                )
            else:
                # Chunk body rows dengan ukuran tertentu dan ulangi header
                for i in range(0, len(body_rows), chunk_size):
                    chunk_slice = body_rows[i : i + chunk_size]
                    table_md = format_table_as_markdown(headers, chunk_slice)
                    blocks.append(
                        Block(
                            rel_path=meta["rel_path"],
                            file_hash=meta["file_hash"],
                            filename=meta["filename"],
                            file_type=meta["file_type"],
                            folder=meta["folder"],
                            modified_at=meta["modified_at"],
                            page=None,
                            sheet=sheet_name,
                            section=f"Sheet: {sheet_name}",
                            kind="table",
                            text=table_md,
                        )
                    )

    finally:
        wb.close()

    return blocks


def _read_text_file(path: Path) -> str:
    """Membaca isi file teks dengan deteksi encoding bertahap."""
    encodings = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
    for enc in encodings:
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, OSError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def extract_txt(path: Path, meta: dict[str, Any], settings: Settings) -> list[Block]:
    """Ekstraksi file teks polos (TXT) dengan pemisahan per paragraf."""
    blocks: list[Block] = []
    raw_text = _read_text_file(path)

    # Split berdasarkan 2 newline atau lebih
    parts = re.split(r"\n\s*\n", raw_text)
    for part in parts:
        clean = part.strip()
        if not clean:
            continue
        blocks.append(
            Block(
                rel_path=meta["rel_path"],
                file_hash=meta["file_hash"],
                filename=meta["filename"],
                file_type=meta["file_type"],
                folder=meta["folder"],
                modified_at=meta["modified_at"],
                page=None,
                sheet=None,
                section=None,
                kind="paragraph",
                text=clean,
            )
        )

    return blocks


def extract_md(path: Path, meta: dict[str, Any], settings: Settings) -> list[Block]:
    """Ekstraksi dokumen Markdown dengan preservasi struktur hierarki heading."""
    blocks: list[Block] = []
    raw_text = _read_text_file(path)
    current_section: Optional[str] = None

    # Pisahkan dokumen menjadi sekuens blok berdasarkan double newline
    parts = re.split(r"\n\s*\n", raw_text)

    for part in parts:
        clean = part.strip()
        if not clean:
            continue

        # Cek jika blok adalah heading markdown (# Heading)
        if clean.startswith("#"):
            first_line = clean.splitlines()[0]
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", first_line)
            if heading_match:
                current_section = heading_match.group(2).strip()
                # Jika seluruh blok hanya baris heading
                if len(clean.splitlines()) == 1:
                    blocks.append(
                        Block(
                            rel_path=meta["rel_path"],
                            file_hash=meta["file_hash"],
                            filename=meta["filename"],
                            file_type=meta["file_type"],
                            folder=meta["folder"],
                            modified_at=meta["modified_at"],
                            page=None,
                            sheet=None,
                            section=current_section,
                            kind="heading",
                            text=clean,
                        )
                    )
                    continue

        # Cek tabel markdown (| a | b |)
        if clean.startswith("|") and "\n|" in clean:
            blocks.append(
                Block(
                    rel_path=meta["rel_path"],
                    file_hash=meta["file_hash"],
                    filename=meta["filename"],
                    file_type=meta["file_type"],
                    folder=meta["folder"],
                    modified_at=meta["modified_at"],
                    page=None,
                    sheet=None,
                    section=current_section,
                    kind="table",
                    text=clean,
                )
            )
            continue

        # Cek code block (```)
        if clean.startswith("```") and clean.endswith("```"):
            blocks.append(
                Block(
                    rel_path=meta["rel_path"],
                    file_hash=meta["file_hash"],
                    filename=meta["filename"],
                    file_type=meta["file_type"],
                    folder=meta["folder"],
                    modified_at=meta["modified_at"],
                    page=None,
                    sheet=None,
                    section=current_section,
                    kind="code",
                    text=clean,
                )
            )
            continue

        # Blok teks paragraf / list
        kind = "list" if clean.startswith(("- ", "* ", "1. ", "2. ")) else "paragraph"
        blocks.append(
            Block(
                rel_path=meta["rel_path"],
                file_hash=meta["file_hash"],
                filename=meta["filename"],
                file_type=meta["file_type"],
                folder=meta["folder"],
                modified_at=meta["modified_at"],
                page=None,
                sheet=None,
                section=current_section,
                kind=kind,
                text=clean,
            )
        )

    return blocks


# ==============================================================================
# Dispatcher & Orchestration
# ==============================================================================


def load_file(
    path: Path,
    root_dir: Optional[Path] = None,
    settings: Optional[Settings] = None,
) -> list[Block]:
    """Memuat dan mengekstrak blok dari satu file dokumen."""
    if settings is None:
        settings = get_settings()

    path = Path(path)
    if not path.exists() or not path.is_file():
        logger.warning(f"File tidak ditemukan: {path}")
        return []

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        logger.warning(f"Format file tidak didukung: {path.name}")
        return []

    try:
        meta = get_file_metadata(path, root_dir)
        if ext == ".pdf":
            return extract_pdf(path, meta, settings)
        if ext == ".docx":
            return extract_docx(path, meta, settings)
        if ext == ".xlsx":
            return extract_xlsx(path, meta, settings)
        if ext == ".md":
            return extract_md(path, meta, settings)
        if ext == ".txt":
            return extract_txt(path, meta, settings)
    except Exception as exc:
        # Isolasi error per file agar tidak menggagalkan seluruh proses
        logger.error(f"Gagal memuat file {path.name} ({exc.__class__.__name__}): {exc}")
        return []

    return []


def load_documents(
    dir_path: Path,
    settings: Optional[Settings] = None,
) -> tuple[list[Block], dict[str, Any]]:
    """Memuat seluruh dokumen dalam direktori dan mengumpulkan statistik."""
    if settings is None:
        settings = get_settings()

    dir_path = Path(dir_path)
    files = scan_directory(dir_path, max_file_mb=settings.MAX_FILE_MB)

    all_blocks: list[Block] = []
    seen_hashes: dict[str, str] = {}  # hash -> first_rel_path
    duplicates: list[dict[str, str]] = []
    failed_files: list[str] = []
    loaded_files: list[str] = []

    for file_path in files:
        try:
            f_hash = compute_file_hash(file_path)
            meta = get_file_metadata(file_path, dir_path)
            rel_path = meta["rel_path"]

            if f_hash in seen_hashes:
                duplicates.append(
                    {
                        "original": seen_hashes[f_hash],
                        "duplicate": rel_path,
                        "file_hash": f_hash,
                    }
                )
                logger.info(f"File duplikat terdeteksi: {rel_path} (identik dengan {seen_hashes[f_hash]})")
            else:
                seen_hashes[f_hash] = rel_path

            blocks = load_file(file_path, root_dir=dir_path, settings=settings)
            if blocks:
                all_blocks.extend(blocks)
                loaded_files.append(rel_path)
            else:
                # File mungkin kosong atau gagal diekstrak
                logger.warning(f"Tidak ada blok yang diekstrak dari {rel_path}")
        except Exception as exc:
            failed_files.append(file_path.name)
            logger.error(f"Kesalahan tak terduga pada {file_path.name}: {exc}")

    # Rekapitulasi statistik
    counts_by_kind: dict[str, int] = {}
    counts_by_type: dict[str, int] = {}
    ocr_pages = 0

    for b in all_blocks:
        counts_by_kind[b.kind] = counts_by_kind.get(b.kind, 0) + 1
        counts_by_type[b.file_type] = counts_by_type.get(b.file_type, 0) + 1
        if b.needs_ocr:
            ocr_pages += 1

    stats: dict[str, Any] = {
        "total_scanned_files": len(files),
        "loaded_files_count": len(loaded_files),
        "failed_files": failed_files,
        "duplicates": duplicates,
        "total_blocks": len(all_blocks),
        "blocks_by_kind": counts_by_kind,
        "blocks_by_type": counts_by_type,
        "needs_ocr_pages": ocr_pages,
        "total_chars": sum(len(b.text) for b in all_blocks),
    }

    return all_blocks, stats


# ==============================================================================
# CLI Entrypoint
# ==============================================================================


def main() -> None:
    """Entry point CLI untuk pengujian mandiri loader dokumen."""
    parser = argparse.ArgumentParser(
        description="Loader dan normalisasi dokumen lokal untuk asisten pribadi."
    )
    parser.add_argument(
        "directory",
        type=str,
        nargs="?",
        default="data/sample",
        help="Direktori yang berisi dokumen (default: data/sample)",
    )
    parser.add_argument(
        "--stats",
        action="store_true",
        help="Tampilkan statistik ringkasan dokumen yang dimuat",
    )
    parser.add_argument(
        "--preview",
        type=int,
        default=0,
        help="Tampilkan cuplikan N blok pertama",
    )

    args = parser.parse_args()
    setup_logging()

    target_dir = Path(args.directory)
    if not target_dir.exists():
        print(f"Error: Direktori '{target_dir}' tidak ditemukan.")
        return

    print("=" * 70)
    print(f"MEMUAT DOKUMEN DARI: {target_dir.resolve()}")
    print("=" * 70)

    blocks, stats = load_documents(target_dir)

    print("\nHASIL PEMINDAIAN DAN EKSTRAKSI:")
    print(f"- Total file ditemukan      : {stats['total_scanned_files']}")
    print(f"- File berhasil dimuat      : {stats['loaded_files_count']}")
    print(f"- File gagal                : {len(stats['failed_files'])}")
    print(f"- File duplikat terdeteksi  : {len(stats['duplicates'])}")
    for dup in stats["duplicates"]:
        print(f"  * Duplikat: {dup['duplicate']} == {dup['original']}")
    print(f"- Total blok teks/tabel     : {stats['total_blocks']}")
    print(f"- Total karakter diekstrak  : {stats['total_chars']:,}")
    print(f"- Halaman butuh OCR         : {stats['needs_ocr_pages']}")

    print("\nDistribusi Format File:")
    for ftype, count in sorted(stats["blocks_by_type"].items()):
        print(f"  * {ftype.upper():<6} : {count} blok")

    print("\nDistribusi Jenis Blok:")
    for kind, count in sorted(stats["blocks_by_kind"].items()):
        print(f"  * {kind:<10} : {count} blok")

    if args.preview > 0:
        preview_count = min(args.preview, len(blocks))
        print("\n" + "=" * 70)
        print(f"CUPLIKAN {preview_count} BLOK PERTAMA")
        print("=" * 70)
        for idx, b in enumerate(blocks[:preview_count], start=1):
            loc = format_location(b)
            print(f"\n[Blok #{idx}] {loc} (kind: {b.kind}, tipe: {b.file_type})")
            short_text = b.text if len(b.text) <= 200 else b.text[:197] + "..."
            print(short_text)

    print("\n" + "=" * 70)
    print("PROSES SELESAI")
    print("=" * 70)


if __name__ == "__main__":
    main()
